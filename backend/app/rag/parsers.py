"""
STAGE i — DATA INGESTION LAYER

Turns whatever the user uploaded into plain text, so the rest of the
pipeline never has to care what format a document originally was.

Supported today: PDF, DOCX, HTML, images (OCR), CSV, XLSX, JSON, Markdown,
and plain text.

To add a new file type:
  1. Write a new `parse_xxx(path) -> str` function
  2. Add its extension(s) to EXTENSION_TO_TYPE
  3. Register the type -> parser mapping in FILE_TYPE_PARSERS
Nothing else in the pipeline needs to change — file-type detection,
dispatch, and every caller all flow through this one file.
"""

import os
import csv
import json


def parse_pdf(path: str) -> str:
    """
    Table-aware PDF parsing: pdfplumber (if installed) extracts each page's
    tables with rows kept intact — critical for spec sheets/catalogs where a
    value (e.g. "5 kWh") is meaningless without its row label (e.g. "PowerCell
    Home 5"). Plain narrative text still comes from pypdf either way.

    Falls back to pypdf-only text extraction if pdfplumber isn't installed.
    """
    try:
        from pypdf import PdfReader
    except ImportError:
        return "[PDF parsing unavailable — install pypdf: pip install pypdf]"

    try:
        reader = PdfReader(path)
        plain_pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as e:
        return f"[Failed to parse PDF: {e}]"

    table_text = ""
    try:
        import pdfplumber
        table_lines = []
        with pdfplumber.open(path) as pdf:
            for page_num, page in enumerate(pdf.pages, start=1):
                for table in page.extract_tables() or []:
                    table_lines.append(f"--- Table on page {page_num} ---")
                    for row in table:
                        cells = [str(c).strip() if c is not None else "" for c in row]
                        if any(cells):
                            table_lines.append(" | ".join(cells))
        table_text = "\n".join(table_lines)
    except ImportError:
        pass
    except Exception:
        pass

    combined = "\n\n".join(plain_pages).strip()
    if table_text:
        combined += "\n\n" + table_text
    return combined.strip()


def parse_docx(path: str) -> str:
    try:
        import docx
    except ImportError:
        return "[DOCX parsing unavailable — install python-docx: pip install python-docx]"

    try:
        doc = docx.Document(path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception as e:
        return f"[Failed to parse DOCX: {e}]"


def parse_html(path: str) -> str:
    try:
        from bs4 import BeautifulSoup
    except ImportError:
        return "[HTML parsing unavailable — install beautifulsoup4: pip install beautifulsoup4]"

    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            soup = BeautifulSoup(f.read(), "html.parser")
        for tag in soup(["script", "style"]):
            tag.decompose()
        text = soup.get_text(separator="\n")
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        return "\n".join(lines)
    except Exception as e:
        return f"[Failed to parse HTML: {e}]"


def parse_image(path: str) -> str:
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return "[OCR unavailable — install pytesseract + Pillow: pip install pytesseract Pillow]"

    try:
        text = pytesseract.image_to_string(Image.open(path))
        return text.strip() or "[OCR ran but found no text in this image]"
    except Exception as e:
        return f"[OCR unavailable — is the Tesseract binary installed on this machine? ({e})]"


def parse_csv(path: str) -> str:
    try:
        lines = []
        with open(path, "r", encoding="utf-8", errors="ignore", newline="") as f:
            reader = csv.reader(f)
            header = next(reader, None)
            if header:
                lines.append(" | ".join(header))
            for row in reader:
                if any(cell.strip() for cell in row):
                    lines.append(" | ".join(row))
        return "\n".join(lines)
    except Exception as e:
        return f"[Failed to parse CSV: {e}]"


def parse_xlsx(path: str) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "[XLSX parsing unavailable — install openpyxl: pip install openpyxl]"

    try:
        wb = load_workbook(path, data_only=True, read_only=True)
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"--- Sheet: {sheet.title} ---")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    except Exception as e:
        return f"[Failed to parse XLSX: {e}]"


def parse_json(path: str) -> str:
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            data = json.load(f)
        return json.dumps(data, indent=2, ensure_ascii=False)
    except Exception as e:
        return f"[Failed to parse JSON: {e}]"


def parse_text(path: str) -> str:
    try:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception as e:
        return f"[Failed to read text file: {e}]"


EXTENSION_TO_TYPE = {
    "pdf": "pdf",
    "docx": "docx",
    "doc": "docx",
    "html": "html",
    "htm": "html",
    "jpg": "image",
    "jpeg": "image",
    "png": "image",
    "bmp": "image",
    "tiff": "image",
    "csv": "csv",
    "xlsx": "xlsx",
    "xls": "xlsx",
    "json": "json",
    "md": "text",
    "markdown": "text",
    "txt": "text",
}

FILE_TYPE_PARSERS = {
    "pdf": parse_pdf,
    "docx": parse_docx,
    "html": parse_html,
    "image": parse_image,
    "csv": parse_csv,
    "xlsx": parse_xlsx,
    "json": parse_json,
    "text": parse_text,
}


def detect_file_type(filename: str) -> str:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    return EXTENSION_TO_TYPE.get(ext, "text")


def parse_document(storage_path: str, file_type: str) -> str:
    parser = FILE_TYPE_PARSERS.get(file_type, parse_text)
    if not os.path.exists(storage_path):
        return f"[File not found on disk: {storage_path}]"
    return parser(storage_path)